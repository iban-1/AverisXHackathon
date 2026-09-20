"""M3/M7: parse an SI/BL attachment (txt/pdf/docx/xlsx) into the 7
canonical fields.

txt and PDF-extracted text render as "Label: value" or, in PDF's
whitespace-column layout, "Label value" with no colon (see synonyms.py for
the label survey this was built from). DOCX/XLSX are table/cell based and
are parsed structurally instead (extract_fields_docx/_xlsx). Address/
description continuation lines that follow a name field are intentionally
not folded into the value — the compared identity is the name on the
label's own line/cell.
"""
import re
from pathlib import Path

from .synonyms import LABEL_TO_FIELD, field_for_label, is_other_doc_type_label

# Label text itself is unconstrained (some documents mix in CJK characters,
# e.g. "Gross Weight毛重(KGS):") — synonyms.normalize_label() strips those
# before dictionary lookup. Only the ":" separator anchors the split.
_LABEL_LINE = re.compile(r"^([^:\n]{1,60}):\s*(.*)$")


def _label_pattern(label: str) -> str:
    return r"\s+".join(re.escape(word) for word in label.split(" "))


# PDF's column layout has no colon ("Shipper APRIL FAR EAST...") — fall
# back to matching one of our known label phrases directly, longest first
# so e.g. "Port of Discharge (POD)" wins over the bare "POD"/"Port of
# Discharge" alternatives.
_NO_COLON_LABEL_RE = re.compile(
    r"^(" + "|".join(_label_pattern(l) for l in sorted(LABEL_TO_FIELD, key=len, reverse=True)) + r")\s+(.+)$",
    re.IGNORECASE,
)

# Placeholder values used by the "missing_value" edge cases (blank fields
# that are uncertainty, not a discrepancy) — extraction treats these as if
# the field were absent so downstream comparison/reliability logic (M6) can
# tell "not present" from "present but blank". Covers bare "???"/"_______",
# a blank with a stray leftover unit ("____MT"), and word placeholders
# (N/A, TBA, NIL, PENDING) — the full set observed across the dataset.
_PLACEHOLDER_RE = re.compile(
    r"^(?:[_?]+\s*(?:MTS?|KGS?)?|N/?A|TBA|NIL|PENDING)$", re.IGNORECASE
)


def _is_usable_value(value: str) -> bool:
    return bool(value) and not _PLACEHOLDER_RE.match(value)


# Some PDFs' text layer interleaves an overlapping label and value (e.g. the
# long "Notify Party/Intermediate Consignee" wraps onto its own value:
# "Notify Party/Intermediate ConsNigAnGeAePPA EXPORTS" in the raw text) —
# the exact-label match then fails and a shorter alternative ("notify")
# matches instead, capturing leftover label fragments as if they were the
# value. A real company/party name is extremely unlikely to contain this
# document's own boilerplate vocabulary, so treat that as corruption rather
# than a real value — leaving the field unset routes to missing_value
# (escalate) instead of a false-positive mismatch.
_LABEL_LEAKAGE_RE = re.compile(
    r"\b(?:notify|consignee|shipper|intermediate)\b", re.IGNORECASE
)
_NAME_FIELDS = {"shipper", "consignee", "notify_party"}


# Some PDFs' text layer glues stray characters between "Weight" and its unit
# suffix (e.g. "TOTAL Gross Weightnn(KGS): 143,940 KG" — a font/kerning
# artifact, not a real label variant) which the exact-label matchers above
# can't recognize. Last-resort: any "gross w(eigh)?t ... : N KG" pattern,
# however it's glued together, unambiguously names this one field.
_GROSS_WEIGHT_FALLBACK_RE = re.compile(
    r"gross\s*w(?:eigh)?t.{0,20}?:\s*([\d,]+(?:\.\d+)?)\s*kgs?\b", re.IGNORECASE
)


def extract_fields_txt(text: str) -> dict[str, str]:
    """Return {field_name: raw_value_string} for the 7 fields found in
    plain-text or PDF-extracted SI/BL text. A field with only a placeholder
    value (`???`, `_______`, `TBA`, `N/A`, ...) or no attachment text at all
    is omitted."""
    fields: dict[str, str] = {}
    for raw_line in text.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        m = _LABEL_LINE.match(line)
        if m:
            label, value = m.group(1), m.group(2).strip()
        else:
            m2 = _NO_COLON_LABEL_RE.match(line)
            if not m2:
                continue
            label, value = m2.group(1), m2.group(2).strip()
        field = field_for_label(label)
        if not field or field in fields:
            continue
        if not _is_usable_value(value):
            continue
        if field in _NAME_FIELDS and _LABEL_LEAKAGE_RE.search(value):
            continue
        fields[field] = value

    if "gross_weight_kg" not in fields:
        m = _GROSS_WEIGHT_FALLBACK_RE.search(text)
        if m:
            fields["gross_weight_kg"] = f"{m.group(1)} KG"

    return fields


def extract_fields_pdf(path) -> dict[str, str]:
    """Flatten every page's text and reuse the same line parser as txt."""
    import pdfplumber

    with pdfplumber.open(path) as pdf:
        text = "\n".join(page.extract_text() or "" for page in pdf.pages)
    return extract_fields_txt(text)


def extract_fields_docx(path) -> dict[str, str]:
    """DOCX renders SI/BL fields as a 2-column table: [label, value]. The
    value cell's first line is the compared identity; later lines are the
    address/description continuation (same convention as txt/PDF)."""
    import docx

    fields: dict[str, str] = {}
    document = docx.Document(path)
    for table in document.tables:
        for row in table.rows:
            if len(row.cells) < 2:
                continue
            label = row.cells[0].text.strip()
            value = row.cells[1].text.strip().splitlines()[0].strip() if row.cells[1].text.strip() else ""
            field = field_for_label(label)
            if not field or field in fields or not _is_usable_value(value):
                continue
            fields[field] = value
    return fields


def extract_fields_xlsx(path) -> dict[str, str]:
    """XLSX renders SI/BL fields as [label, value] rows (one or more
    sheets). Multi-part values are "|"-joined ("NAME | ADDRESS LINE"); the
    part before the first "|" is the compared identity. Weight is
    sometimes a raw numeric cell rather than a "N KG" string."""
    import openpyxl

    fields: dict[str, str] = {}
    workbook = openpyxl.load_workbook(path, data_only=True)
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows(values_only=True):
            if len(row) < 2 or row[0] is None or row[1] is None:
                continue
            label = str(row[0]).strip()
            raw_value = row[1]
            value = raw_value.split("|")[0].strip() if isinstance(raw_value, str) else str(raw_value)
            field = field_for_label(label)
            if not field or field in fields or not _is_usable_value(value):
                continue
            fields[field] = value
    return fields


_EXTRACTORS_BY_SUFFIX = {
    ".txt": lambda p: extract_fields_txt(Path(p).read_text(encoding="utf-8")),
    ".pdf": extract_fields_pdf,
    ".docx": extract_fields_docx,
    ".xlsx": extract_fields_xlsx,
}


def extract_fields(path) -> dict[str, str]:
    """Dispatch to the right extractor by file extension."""
    suffix = Path(path).suffix.lower()
    extractor = _EXTRACTORS_BY_SUFFIX.get(suffix)
    if extractor is None:
        raise ValueError(f"no extractor for {suffix!r} ({path})")
    return extractor(path)


_EXPECTED_HEADERS = ("SHIPPING INSTRUCTION", "BILL OF LADING")


def looks_like_other_doc_type(text: str) -> bool:
    """True if a "BL"/"SI" attachment is actually some other document type
    (Commercial Invoice, Packing List, Certificate of Origin). Primary
    signal: the document's own header line doesn't name it as either an SI
    or a BL. Secondary/defense-in-depth signal: labels from that other
    document family (invoice/certificate fields — see
    synonyms.OTHER_DOC_TYPE_LABELS) rather than any of the 7 compared ones."""
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    if lines and not any(h in lines[0].upper() for h in _EXPECTED_HEADERS):
        return True
    for line in text.splitlines():
        m = _LABEL_LINE.match(line)
        if m and is_other_doc_type_label(m.group(1)):
            return True
    return False


def normalize_text_value(value: str) -> str:
    """Case/whitespace-insensitive form for comparing name/port fields."""
    return " ".join(value.strip().split()).upper()


_LEADING_INT_RE = re.compile(r"(\d+)")


def parse_container_count(value: str) -> int | None:
    """'1 x 40'HC' -> 1"""
    m = _LEADING_INT_RE.search(value)
    return int(m.group(1)) if m else None


_NUMBER_RE = re.compile(r"([\d,]+(?:\.\d+)?)")


def parse_gross_weight_kg(value: str) -> float | None:
    """'21,577 KG' -> 21577.0"""
    m = _NUMBER_RE.search(value)
    return float(m.group(1).replace(",", "")) if m else None
